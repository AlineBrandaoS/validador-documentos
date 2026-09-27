"""
VALIDADOR DE LEGIBILIDADE DE DOCUMENTOS (mobilização / contratação)
====================================================================

Analisa PDFs e imagens de documentos e classifica cada um como:
    LEGÍVEL   -> pode seguir no processo
    ILEGÍVEL  -> pedir novo envio ao colaborador (com o motivo)
    REVISAR   -> caso duvidoso: uma pessoa da equipe confere

Etapas de cada página:
    1. Filtro técnico (OpenCV + Tesseract): resolução, página em branco,
       desfoque, brilho, contraste e confiança do OCR.
    2. Análise da IA (Claude, com visão): legível/ilegível, tipo do
       documento, problemas encontrados e nível de certeza.
    3. Decisão: só aprova ou reprova automaticamente quando há segurança.
       Qualquer dúvida ou divergência entre IA e filtro técnico -> REVISAR.

Como usar (no terminal, dentro da pasta do programa):
    python validador_legibilidade.py "C:\\Documentos\\Mobilizacao"
    python validador_legibilidade.py "C:\\Documentos\\Mobilizacao" --sem-ia
    python validador_legibilidade.py --calibrar "C:\\Documentos\\teste_calibracao"

Organização recomendada da pasta: uma subpasta por colaborador.
    Mobilizacao\\
        JOAO DA SILVA\\rg.pdf, cnh.jpg, aso.pdf ...
        MARIA SOUZA\\...
"""

from __future__ import annotations

import argparse
import base64
import hashlib
import io
import json
import os
import sys
import threading
import unicodedata
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path
from typing import Any

import cv2
import numpy as np
import pandas as pd
from PIL import Image, ImageOps, ImageSequence

try:
    from dotenv import load_dotenv
    load_dotenv(Path(__file__).with_name(".env"))
except ImportError:  # python-dotenv é opcional
    pass

try:
    import pytesseract
except ImportError:
    pytesseract = None

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import anthropic
except ImportError:
    anthropic = None


# =============================================================================
# CONFIGURAÇÕES  (ajuste aqui ou no arquivo .env)
# =============================================================================

MODELO_IA = os.getenv("MODELO_IA", "claude-sonnet-5")
CONFIANCA_MINIMA = float(os.getenv("CONFIANCA_MINIMA", "0.85"))  # abaixo disso -> REVISAR
THREADS = int(os.getenv("THREADS", "4"))                         # arquivos analisados em paralelo
DPI_PDF = 200                 # resolução usada para transformar PDF em imagem
MAX_PAGINAS_POR_ARQUIVO = 20  # acima disso o arquivo vai para REVISAR
LADO_MAX_IA = 1568            # tamanho máximo (px) da imagem enviada à IA

# Limites do filtro técnico. Ajuste depois da calibração (--calibrar).
LIMITES = {
    "resolucao_minima": 300,      # menor lado em px; abaixo -> ILEGÍVEL direto
    "contraste_em_branco": 8.0,   # página em branco = contraste abaixo disso
    "bordas_em_branco": 0.002,    #   E densidade de bordas abaixo disso (as duas condições)
    "desfoque_alerta": 40.0,      # variância do Laplaciano; abaixo -> alerta de desfoque
    "brilho_min": 50.0,           # média de cinza; abaixo -> muito escuro
    "contraste_alerta": 15.0,     # abaixo -> baixo contraste
    "ocr_palavras_min": 8,        # só avalia a confiança do OCR com pelo menos N palavras
    "ocr_confianca_alerta": 35.0, # confiança média do Tesseract (0-100); abaixo -> alerta
}

EXTENSOES_PDF = {".pdf"}
EXTENSOES_IMAGEM = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".bmp", ".webp"}

TIPOS_DOCUMENTO = [
    "RG", "CPF", "CNH", "CTPS", "TITULO_ELEITOR", "CERTIFICADO_RESERVISTA",
    "COMPROVANTE_RESIDENCIA", "CERTIDAO_NASCIMENTO_OU_CASAMENTO",
    "CERTIDAO_NASCIMENTO_FILHO", "CARTAO_PIS_NIS", "ASO", "CERTIFICADO_NR",
    "COMPROVANTE_ESCOLARIDADE", "CARTAO_VACINA", "DADOS_BANCARIOS",
    "FOTO_3X4", "FICHA_REGISTRO_OU_CONTRATO", "FICHA_EPI", "ORDEM_SERVICO",
    "CERTIFICADO_RAC", "PRO_RAC", "CONTROLE_FREQUENCIA_ART", "FREQUENCIA_PST",
    "CERTIFICADO_TREINAMENTO_EQUIPAMENTO", "CARTA_ANUENCIA", "ATESTADO_PROFICIENCIA",
    "CONTEUDO_PROGRAMATICO", "OUTRO_DOCUMENTO", "NAO_E_DOCUMENTO",
]

PROBLEMAS = [
    "BORRADO_OU_DESFOCADO", "ESCURO", "REFLEXO_OU_ESTOURADO", "CORTADO_OU_INCOMPLETO",
    "BAIXA_RESOLUCAO", "DADOS_COBERTOS_OU_RASURADOS", "TORTO_OU_DISTORCIDO",
    "PAGINA_EM_BRANCO", "NAO_E_DOCUMENTO", "OUTRO",
]

# Mude este número sempre que alterar o prompt: invalida o cache antigo.
VERSAO_PROMPT = "2"

LEGIVEL, ILEGIVEL, REVISAR, EM_BRANCO = "LEGÍVEL", "ILEGÍVEL", "REVISAR", "EM BRANCO"


# =============================================================================
# 1. LEITURA DOS ARQUIVOS
# =============================================================================

_pdf_lock = threading.Lock()


def carregar_paginas(caminho: Path) -> list[Image.Image]:
    """Devolve a lista de páginas do arquivo como imagens RGB."""
    ext = caminho.suffix.lower()

    if ext in EXTENSOES_PDF:
        if fitz is None:
            raise RuntimeError("PyMuPDF não instalado (pip install pymupdf)")
        paginas = []
        with _pdf_lock, fitz.open(caminho) as doc:  # PyMuPDF não é seguro entre threads
            if doc.needs_pass:
                raise RuntimeError("PDF protegido por senha")
            if doc.page_count == 0:
                raise RuntimeError("PDF sem páginas")
            if doc.page_count > MAX_PAGINAS_POR_ARQUIVO:
                raise RuntimeError(
                    f"PDF com {doc.page_count} páginas (limite {MAX_PAGINAS_POR_ARQUIVO})")
            for pagina in doc:
                pix = pagina.get_pixmap(dpi=DPI_PDF, alpha=False)
                paginas.append(Image.frombytes("RGB", (pix.width, pix.height), pix.samples))
        return paginas

    if ext in EXTENSOES_IMAGEM:
        with Image.open(caminho) as img:
            paginas = []
            for quadro in ImageSequence.Iterator(img):  # TIFF pode ter várias páginas
                quadro = ImageOps.exif_transpose(quadro)  # corrige foto girada do celular
                paginas.append(quadro.convert("RGB"))
                if len(paginas) > MAX_PAGINAS_POR_ARQUIVO:
                    raise RuntimeError("Imagem com páginas demais")
        return paginas

    raise RuntimeError(f"Formato não suportado: {ext}")


# =============================================================================
# 2. FILTRO TÉCNICO (OpenCV + Tesseract)
# =============================================================================

_ocr_lock = threading.Lock()
_ocr_estado: dict = {"verificado": False, "idioma": None, "aviso": None}


def _preparar_tesseract() -> str | None:
    """Descobre se o Tesseract está instalado e qual idioma usar (por > eng)."""
    with _ocr_lock:
        if _ocr_estado["verificado"]:
            return _ocr_estado["idioma"]
        _ocr_estado["verificado"] = True
        if pytesseract is None:
            _ocr_estado["aviso"] = "pytesseract não instalado: OCR desativado"
            return None
        cmd = os.getenv("TESSERACT_CMD")
        padrao_windows = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")
        if cmd:
            pytesseract.pytesseract.tesseract_cmd = cmd
        elif padrao_windows.exists():
            pytesseract.pytesseract.tesseract_cmd = str(padrao_windows)
        try:
            idiomas = set(pytesseract.get_languages(config=""))
        except Exception:
            _ocr_estado["aviso"] = "Tesseract não encontrado: OCR desativado (confira TESSERACT_CMD no .env)"
            return None
        if "por" in idiomas:
            _ocr_estado["idioma"] = "por"
        else:
            _ocr_estado["idioma"] = "eng"
            _ocr_estado["aviso"] = "Idioma 'por' do Tesseract não instalado: usando 'eng'"
        return _ocr_estado["idioma"]


def _redimensionar(img_array: np.ndarray, lado_max: int) -> np.ndarray:
    h, w = img_array.shape[:2]
    escala = lado_max / max(h, w)
    if escala >= 1:
        return img_array
    return cv2.resize(img_array, (int(w * escala), int(h * escala)), interpolation=cv2.INTER_AREA)


def analisar_tecnico(img: Image.Image) -> dict:
    """Mede a qualidade da imagem. Não decide sozinho, exceto em casos extremos."""
    cinza_original = np.array(img.convert("L"))
    altura, largura = cinza_original.shape
    menor_lado = min(altura, largura)

    # Normaliza o tamanho para os números serem comparáveis entre fotos e scans
    cinza = _redimensionar(cinza_original, 1200)
    desfoque = float(cv2.Laplacian(cinza, cv2.CV_64F).var())
    brilho = float(cinza.mean())
    contraste = float(cinza.std())
    bordas = cv2.Canny(cinza, 50, 150)
    fracao_bordas = float((bordas > 0).mean())

    resultado = {
        "largura": largura, "altura": altura,
        "desfoque": round(desfoque, 1), "brilho": round(brilho, 1),
        "contraste": round(contraste, 1), "bordas": round(fracao_bordas, 4),
        "ocr_palavras": None, "ocr_confianca": None,
        "bloqueio": None,   # motivo de reprovação técnica direta
        "em_branco": False,
        "alertas": [],      # sinais de problema que a IA precisa confirmar
    }

    if menor_lado < LIMITES["resolucao_minima"]:
        resultado["bloqueio"] = f"Resolução muito baixa ({largura}x{altura}px)"
        return resultado

    # As DUAS condições juntas: documento borrado ou escuro não pode ser confundido com página vazia
    if contraste < LIMITES["contraste_em_branco"] and fracao_bordas < LIMITES["bordas_em_branco"]:
        resultado["em_branco"] = True
        return resultado

    if desfoque < LIMITES["desfoque_alerta"]:
        resultado["alertas"].append(f"possível desfoque (nitidez {desfoque:.0f})")
    if brilho < LIMITES["brilho_min"]:
        resultado["alertas"].append(f"imagem escura (brilho {brilho:.0f})")
    if contraste < LIMITES["contraste_alerta"]:
        resultado["alertas"].append(f"baixo contraste ({contraste:.0f})")

    idioma = _preparar_tesseract()
    if idioma and pytesseract is not None:
        try:
            img_ocr = Image.fromarray(_redimensionar(np.array(img), 2200))
            dados = pytesseract.image_to_data(
                img_ocr, lang=idioma, output_type=pytesseract.Output.DICT, timeout=60)
            confiancas = [
                float(c) for c, t in zip(dados["conf"], dados["text"])
                if str(t).strip() and len(str(t).strip()) >= 2 and float(c) >= 0
            ]
            resultado["ocr_palavras"] = len(confiancas)
            if confiancas:
                media = sum(confiancas) / len(confiancas)
                resultado["ocr_confianca"] = round(media, 1)
                if (len(confiancas) >= LIMITES["ocr_palavras_min"]
                        and media < LIMITES["ocr_confianca_alerta"]):
                    resultado["alertas"].append(f"OCR com baixa confiança ({media:.0f}%)")
        except Exception as erro:  # OCR é apoio: se falhar, segue sem ele
            resultado["alertas"].append(f"OCR falhou ({type(erro).__name__})")

    return resultado


# =============================================================================
# 3. ANÁLISE COM IA (Claude)
# =============================================================================

PROMPT_SISTEMA = """Você é um analista de Departamento Pessoal que confere documentos \
enviados por colaboradores para admissão e mobilização em obra.

Sua única tarefa é avaliar se a imagem está LEGÍVEL o suficiente para o RH conferir \
e cadastrar os dados sem precisar pedir um novo envio. Não avalie autenticidade, \
validade ou se os dados estão corretos.

Critérios:
- LEGÍVEL: todos os dados essenciais que aparecem no documento podem ser lidos com \
segurança, sem adivinhar nenhum caractere: nome, números (CPF, RG, CNH, PIS, CTPS \
etc.), datas, órgão emissor e, quando existirem, foto, assinatura, carimbo e QR code.
- ILEGÍVEL: qualquer dado essencial borrado, desfocado, cortado pela borda da \
imagem, coberto (dedo, reflexo, sombra, rasura, dobra), escuro ou claro demais, \
em resolução baixa demais, ou documento visivelmente incompleto.
- Se um único número importante tiver algum dígito duvidoso, o documento é ILEGÍVEL.
- Uma página só com o verso sem dados relevantes pode ser LEGÍVEL (nada a ler), \
indique isso na observação.

O campo "confianca" é a sua certeza sobre a decisão legível/ilegível, de 0 a 1. \
Seja rigoroso: use valores acima de 0.9 só quando não houver dúvida nenhuma. \
Na dúvida, reduza a confiança; o sistema encaminha para revisão humana.

Identifique também o tipo do documento. Para certificado de NR use CERTIFICADO_NR e informe o número da NR no campo "numero" (ex.: "35"). Para RAC (Requisitos de Atividades Críticas, RAC 01 a RAC 09) use CERTIFICADO_RAC, ou PRO_RAC se for o procedimento/PRO da RAC, e informe o número (ex.: "01"). Para os outros tipos, deixe "numero" vazio.

Não transcreva dados pessoais (números, nomes) na observação. Escreva em português, \
de forma curta e objetiva, dizendo o que o RH deve pedir ao colaborador se for \
ilegível. Responda sempre usando a ferramenta registrar_avaliacao."""

FERRAMENTA_IA = {
    "name": "registrar_avaliacao",
    "description": "Registra a avaliação de legibilidade de uma página de documento.",
    "input_schema": {
        "type": "object",
        "properties": {
            "tipo_documento": {"type": "string", "enum": TIPOS_DOCUMENTO},
            "legivel": {"type": "boolean"},
            "confianca": {"type": "number", "minimum": 0, "maximum": 1},
            "problemas": {"type": "array", "items": {"type": "string", "enum": PROBLEMAS}},
            "campos_ilegiveis": {
                "type": "array", "items": {"type": "string"},
                "description": "Nomes dos campos que não dá para ler (ex.: 'número do CPF', 'data de validade').",
            },
            "observacao": {"type": "string", "description": "Frase curta para o RH."},
            "numero": {"type": "string",
                       "description": "Número da NR ou RAC (ex.: '35', '01'); vazio nos outros tipos."},
        },
        "required": ["tipo_documento", "legivel", "confianca", "problemas",
                     "campos_ilegiveis", "observacao"],
    },
}

_cliente_ia = None
_cliente_lock = threading.Lock()


def obter_cliente_ia():
    global _cliente_ia
    with _cliente_lock:
        if _cliente_ia is None:
            if anthropic is None:
                raise RuntimeError("Biblioteca 'anthropic' não instalada (pip install anthropic)")
            if not os.getenv("ANTHROPIC_API_KEY"):
                raise RuntimeError("ANTHROPIC_API_KEY não configurada no arquivo .env")
            _cliente_ia = anthropic.Anthropic(max_retries=5, timeout=180)
        return _cliente_ia


def imagem_para_base64(img: Image.Image) -> str:
    arr = _redimensionar(np.array(img), LADO_MAX_IA)
    buffer = io.BytesIO()
    Image.fromarray(arr).save(buffer, format="JPEG", quality=92)
    return base64.standard_b64encode(buffer.getvalue()).decode("ascii")


def avaliar_com_ia(img: Image.Image, nome_arquivo: str, pagina: int, total: int) -> dict:
    cliente = obter_cliente_ia()
    # Tipados como Any para o VS Code (Pylance) não reclamar dos tipos internos do SDK
    ferramentas: Any = [FERRAMENTA_IA]
    escolha: Any = {"type": "tool", "name": "registrar_avaliacao"}
    mensagens: Any = [{
        "role": "user",
        "content": [
            {"type": "image", "source": {
                "type": "base64", "media_type": "image/jpeg",
                "data": imagem_para_base64(img)}},
            {"type": "text", "text":
                f"Arquivo: {nome_arquivo} (página {pagina} de {total}). "
                "Avalie a legibilidade desta página."},
        ],
    }]
    resposta = cliente.messages.create(
        model=MODELO_IA,
        max_tokens=1024,
        system=PROMPT_SISTEMA,
        tools=ferramentas,
        tool_choice=escolha,
        messages=mensagens,
    )
    for bloco in resposta.content:
        if (getattr(bloco, "type", None) == "tool_use"
                and getattr(bloco, "name", None) == "registrar_avaliacao"):
            return validar_resposta_ia(getattr(bloco, "input", None))
    raise RuntimeError("A IA não devolveu a avaliação no formato esperado")


def validar_resposta_ia(dados: Any) -> dict:
    """Garante que a resposta tem todos os campos e tipos certos."""
    if not isinstance(dados, dict) or not isinstance(dados.get("legivel"), bool):
        raise RuntimeError("Resposta da IA sem o campo 'legivel'")
    try:
        confianca = float(dados["confianca"])
    except (KeyError, TypeError, ValueError):
        raise RuntimeError("Resposta da IA sem o campo 'confianca'")
    tipo = dados.get("tipo_documento")
    return {
        "tipo_documento": tipo if tipo in TIPOS_DOCUMENTO else "OUTRO_DOCUMENTO",
        "legivel": dados["legivel"],
        "confianca": min(max(confianca, 0.0), 1.0),
        "problemas": [p for p in dados.get("problemas") or [] if p in PROBLEMAS],
        "campos_ilegiveis": [str(c) for c in dados.get("campos_ilegiveis") or []],
        "observacao": str(dados.get("observacao") or "").strip(),
        "numero": "".join(ch for ch in str(dados.get("numero") or "") if ch.isdigit())[:2],
    }


# =============================================================================
# 4. CACHE (não pagar duas vezes pelo mesmo arquivo)
# =============================================================================

class Cache:
    def __init__(self, caminho: Path):
        self.caminho = caminho
        self.lock = threading.Lock()
        try:
            self.dados = json.loads(caminho.read_text(encoding="utf-8"))
        except (FileNotFoundError, json.JSONDecodeError):
            self.dados = {}

    @staticmethod
    def chave(arquivo: Path) -> str:
        h = hashlib.sha256()
        with open(arquivo, "rb") as f:
            for bloco in iter(lambda: f.read(1 << 20), b""):
                h.update(bloco)
        return f"{h.hexdigest()}|{MODELO_IA}|v{VERSAO_PROMPT}"

    def obter(self, chave: str):
        with self.lock:
            return self.dados.get(chave)

    def salvar(self, chave: str, valor) -> None:
        with self.lock:
            self.dados[chave] = valor
            temporario = self.caminho.with_suffix(".tmp")
            temporario.write_text(json.dumps(self.dados, ensure_ascii=False), encoding="utf-8")
            temporario.replace(self.caminho)


# =============================================================================
# 5. DECISÃO
# =============================================================================

TEXTO_PROBLEMA = {
    "BORRADO_OU_DESFOCADO": "borrado/desfocado", "ESCURO": "escuro",
    "REFLEXO_OU_ESTOURADO": "reflexo/claro demais", "CORTADO_OU_INCOMPLETO": "cortado/incompleto",
    "BAIXA_RESOLUCAO": "resolução baixa", "DADOS_COBERTOS_OU_RASURADOS": "dados cobertos/rasurados",
    "TORTO_OU_DISTORCIDO": "torto/distorcido", "PAGINA_EM_BRANCO": "página em branco",
    "NAO_E_DOCUMENTO": "não é documento", "OUTRO": "outro problema",
}


def decidir_pagina(tec: dict, ia: dict | None, erro_ia: str | None) -> tuple[str, str]:
    """Regra de decisão. Na dúvida, sempre REVISAR."""
    if tec["bloqueio"]:
        return ILEGIVEL, tec["bloqueio"]
    if tec["em_branco"]:
        return EM_BRANCO, "Página em branco"
    if erro_ia:
        return REVISAR, f"IA indisponível: {erro_ia}"
    if ia is None:
        alerta = f" | Alertas: {'; '.join(tec['alertas'])}" if tec["alertas"] else ""
        return REVISAR, "Sem análise de IA" + alerta

    if ia["tipo_documento"] == "NAO_E_DOCUMENTO" or "NAO_E_DOCUMENTO" in ia["problemas"]:
        return REVISAR, "Não parece ser um documento. " + ia["observacao"]

    if ia["confianca"] < CONFIANCA_MINIMA:
        situacao = "legível" if ia["legivel"] else "ilegível"
        return REVISAR, (f"IA em dúvida (acha {situacao}, certeza {ia['confianca']:.0%}). "
                         + ia["observacao"])

    if ia["legivel"]:
        if tec["alertas"]:
            return REVISAR, ("IA considerou legível, mas a análise técnica indica: "
                             + "; ".join(tec["alertas"]))
        return LEGIVEL, ia["observacao"] or "OK"

    problemas = ", ".join(str(TEXTO_PROBLEMA.get(p, p)) for p in ia["problemas"]) or "ilegível"
    campos = f" Campos: {', '.join(ia['campos_ilegiveis'])}." if ia["campos_ilegiveis"] else ""
    return ILEGIVEL, f"{problemas}.{campos} {ia['observacao']}".strip()


def consolidar_arquivo(paginas: list[dict]) -> tuple[str, str]:
    """Status do arquivo = o pior status entre as páginas (ignorando páginas em branco)."""
    uteis = [p for p in paginas if p["status"] != EM_BRANCO]
    if not uteis:
        return ILEGIVEL, "Todas as páginas estão em branco"
    for status in (ILEGIVEL, REVISAR):
        com_status = [p for p in uteis if p["status"] == status]
        if com_status:
            if len(paginas) == 1:
                return status, com_status[0]["motivo"]
            return status, " | ".join(f"Pág. {p['pagina']}: {p['motivo']}" for p in com_status)
    if len(paginas) == 1:
        return LEGIVEL, uteis[0]["motivo"]
    return LEGIVEL, f"{len(uteis)} página(s) legível(is)"


# =============================================================================
# 6. PROCESSAMENTO DE UM ARQUIVO
# =============================================================================

def processar_arquivo(arquivo: Path, raiz: Path, usar_ia: bool, cache: Cache | None) -> dict:
    relativo = arquivo.relative_to(raiz)
    if len(relativo.parts) > 1:
        colaborador = relativo.parts[0]
    elif not any(p.is_dir() for p in raiz.iterdir()):
        # Escolheu direto a pasta de UM colaborador (sem subpastas): a pasta é o colaborador
        colaborador = raiz.name
    else:
        colaborador = "-"  # arquivo solto na pasta geral: não é de ninguém
    base = {"Colaborador": colaborador, "Arquivo": arquivo.name, "Caminho": str(relativo)}

    try:
        imagens = carregar_paginas(arquivo)
    except Exception as erro:
        return {**base, "Páginas": 0, "Tipo de documento": "-", "Status": REVISAR,
                "Motivo": f"Não foi possível abrir o arquivo: {erro}", "Certeza IA": None,
                "_paginas": []}

    chave = Cache.chave(arquivo) if (cache and usar_ia) else None
    ia_cache = cache.obter(chave) if (cache is not None and chave) else None
    ia_resultados: list = []
    houve_erro_ia = False

    paginas = []
    for i, img in enumerate(imagens, start=1):
        tec = analisar_tecnico(img)
        ia, erro_ia = None, None
        precisa_ia = usar_ia and not tec["bloqueio"] and not tec["em_branco"]
        if precisa_ia:
            if ia_cache and i <= len(ia_cache) and ia_cache[i - 1]:
                ia = ia_cache[i - 1]
            else:
                try:
                    ia = avaliar_com_ia(img, arquivo.name, i, len(imagens))
                except Exception as erro:
                    erro_ia = f"{type(erro).__name__}: {erro}"[:200]
                    houve_erro_ia = True
        ia_resultados.append(ia)

        status, motivo = decidir_pagina(tec, ia, erro_ia)
        paginas.append({
            **base, "pagina": i, "status": status, "motivo": motivo,
            "tipo": ia["tipo_documento"] if ia else "-",
            "numero": ia.get("numero", "") if ia else "",
            "ia_legivel": ia["legivel"] if ia else None,
            "ia_confianca": ia["confianca"] if ia else None,
            "ia_problemas": ", ".join(ia["problemas"]) if ia else "",
            **{k: tec[k] for k in ("largura", "altura", "desfoque", "brilho", "contraste",
                                   "bordas", "ocr_palavras", "ocr_confianca")},
            "alertas_tecnicos": "; ".join(tec["alertas"]),
        })

    # Só guarda no cache se nenhuma chamada à IA falhou
    if cache is not None and chave and not houve_erro_ia:
        cache.salvar(chave, ia_resultados)

    status, motivo = consolidar_arquivo(paginas)
    tipos = [f"{p['tipo']} {p.get('numero') or ''}".strip()
             for p in paginas if p["tipo"] not in ("-", "OUTRO_DOCUMENTO")]
    tipo = Counter(tipos).most_common(1)[0][0] if tipos else (
        "OUTRO_DOCUMENTO" if any(p["tipo"] == "OUTRO_DOCUMENTO" for p in paginas) else "-")
    certezas = [p["ia_confianca"] for p in paginas if p["ia_confianca"] is not None]

    return {**base, "Páginas": len(paginas), "Tipo de documento": tipo, "Status": status,
            "Motivo": motivo, "Certeza IA": min(certezas) if certezas else None,
            "_paginas": paginas}


def listar_arquivos(raiz: Path) -> list[Path]:
    validas = EXTENSOES_PDF | EXTENSOES_IMAGEM
    return sorted(
        p for p in raiz.rglob("*")
        if p.is_file() and p.suffix.lower() in validas
        and not p.name.startswith(("~$", ".", "_cache_validador", "resultado_legibilidade"))
    )


def processar_pasta(raiz: Path, usar_ia: bool, threads: int,
                    log=print, progresso=None) -> list[dict]:
    """log(texto) recebe as mensagens; progresso(n, total, resultado) é chamado a cada arquivo."""
    arquivos = listar_arquivos(raiz)
    if not arquivos:
        log("Nenhum PDF ou imagem encontrado na pasta.")
        return []

    _preparar_tesseract()
    if _ocr_estado["aviso"]:
        log(f"AVISO: {_ocr_estado['aviso']}")
    if usar_ia:
        obter_cliente_ia()  # falha logo no início se a chave não estiver configurada
        log(f"IA: {MODELO_IA} | certeza mínima: {CONFIANCA_MINIMA:.0%}")
    else:
        log("Modo SEM IA: só filtro técnico (o que não for reprovado vai para REVISAR).")

    cache = Cache(raiz / "_cache_validador.json") if usar_ia else None
    log(f"{len(arquivos)} arquivo(s) encontrado(s).\n")

    resultados = []
    with ThreadPoolExecutor(max_workers=max(1, threads)) as executor:
        futuros = {executor.submit(processar_arquivo, a, raiz, usar_ia, cache): a for a in arquivos}
        for n, futuro in enumerate(as_completed(futuros), start=1):
            r = futuro.result()
            resultados.append(r)
            log(f"[{n}/{len(arquivos)}] {r['Status']:<9} {r['Caminho']}")
            if progresso:
                progresso(n, len(arquivos), r)

    resultados.sort(key=lambda r: (r["Colaborador"], r["Caminho"]))
    return resultados


def executar_analise(raiz: Path, usar_ia: bool = True, threads: int = THREADS,
                     log=print, progresso=None,
                     planilha_colaboradores: Path | None = None) -> tuple[Path | None, Counter]:
    """Analisa a pasta, salva a planilha e devolve (caminho da planilha, contagem por status).
    Se receber a planilha de colaboradores (nome e função), também confere o checklist
    de documentos de admissão e segurança (NRs e RACs) e lista o que está faltando."""
    resultados = processar_pasta(raiz, usar_ia, threads, log, progresso)
    if not resultados and planilha_colaboradores is None:
        return None, Counter()
    destino = raiz / f"resultado_legibilidade_{datetime.now():%Y-%m-%d_%H%M%S}.xlsx"
    salvar_planilha(resultados, destino)
    if planilha_colaboradores is not None:
        import checklist_mobilizacao  # módulo do checklist (mesma pasta)
        checklist_mobilizacao.gerar_checklist(
            resultados, raiz, Path(planilha_colaboradores), destino, usar_ia, log)
    contagem = Counter(r["Status"] for r in resultados)
    log(f"\nLegíveis: {contagem[LEGIVEL]} | Ilegíveis: {contagem[ILEGIVEL]} "
        f"| Revisar: {contagem[REVISAR]}")
    log(f"Planilha salva em: {destino}")
    return destino, contagem


# =============================================================================
# 7. PLANILHA DE RESULTADO
# =============================================================================

CORES_STATUS = {LEGIVEL: "C6EFCE", ILEGIVEL: "FFC7CE", REVISAR: "FFEB9C", EM_BRANCO: "D9D9D9"}


def salvar_planilha(resultados: list[dict], destino: Path, extras: dict | None = None) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    docs = pd.DataFrame([{k: v for k, v in r.items() if k != "_paginas"} for r in resultados])
    if "Certeza IA" in docs:
        docs["Certeza IA"] = docs["Certeza IA"].apply(lambda v: None if pd.isna(v) else round(v, 2))

    paginas = pd.DataFrame([p for r in resultados for p in r["_paginas"]])
    if not paginas.empty:
        paginas = paginas.rename(columns={
            "pagina": "Página", "status": "Status", "motivo": "Motivo", "tipo": "Tipo",
            "ia_legivel": "IA: legível?", "ia_confianca": "IA: certeza",
            "ia_problemas": "IA: problemas", "largura": "Largura px", "altura": "Altura px",
            "desfoque": "Nitidez (Laplaciano)", "brilho": "Brilho", "contraste": "Contraste",
            "bordas": "Densidade de bordas", "ocr_palavras": "OCR: palavras",
            "ocr_confianca": "OCR: confiança %", "alertas_tecnicos": "Alertas técnicos"})

    contagem = docs["Status"].value_counts() if not docs.empty else pd.Series(dtype=int)
    resumo = pd.DataFrame(
        [("Data da análise", datetime.now().strftime("%d/%m/%Y %H:%M")),
         ("Modelo de IA", MODELO_IA),
         ("Total de arquivos", len(docs)),
         (LEGIVEL, int(contagem.get(LEGIVEL, 0))),
         (ILEGIVEL, int(contagem.get(ILEGIVEL, 0))),
         (REVISAR, int(contagem.get(REVISAR, 0)))]
        + list((extras or {}).items()),
        columns=["Indicador", "Valor"])

    with pd.ExcelWriter(destino, engine="openpyxl") as writer:
        resumo.to_excel(writer, sheet_name="Resumo", index=False)
        docs.to_excel(writer, sheet_name="Documentos", index=False)
        if not paginas.empty:
            paginas.to_excel(writer, sheet_name="Detalhe por página", index=False)

        for aba in writer.book.worksheets:
            for celula in aba[1]:
                celula.font = Font(bold=True, color="FFFFFF")
                celula.fill = PatternFill("solid", fgColor="1F4E78")
            aba.freeze_panes = "A2"
            if aba.max_row > 1:
                aba.auto_filter.ref = aba.dimensions
            cabecalho = [c.value for c in aba[1]]
            for idx, nome in enumerate(cabecalho, start=1):
                largura = 70 if nome == "Motivo" else min(
                    max(len(str(nome)), *(len(str(c.value or "")) for c in aba[get_column_letter(idx)])) + 2, 45)
                aba.column_dimensions[get_column_letter(idx)].width = largura
            if "Status" in cabecalho:
                col = cabecalho.index("Status") + 1
                for linha in range(2, aba.max_row + 1):
                    celula = aba.cell(row=linha, column=col)
                    cor = CORES_STATUS.get(str(celula.value))
                    if cor:
                        celula.fill = PatternFill("solid", fgColor=cor)
                        celula.font = Font(bold=True)
            if "Motivo" in cabecalho:
                col = cabecalho.index("Motivo") + 1
                for linha in range(2, aba.max_row + 1):
                    aba.cell(row=linha, column=col).alignment = Alignment(wrap_text=True, vertical="top")


# =============================================================================
# 8. CALIBRAÇÃO (medir a taxa de acerto com documentos de resultado conhecido)
# =============================================================================

def _normalizar(texto: str) -> str:
    texto = unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()
    return texto.strip().upper()


def calibrar(pasta: Path, usar_ia: bool, threads: int) -> None:
    """pasta/legivel/... e pasta/ilegivel/... com documentos que você já conferiu."""
    subpastas = {_normalizar(p.name): p.name for p in pasta.iterdir() if p.is_dir()}
    if "LEGIVEL" not in subpastas or "ILEGIVEL" not in subpastas:
        sys.exit("A pasta de calibração precisa ter as subpastas 'legivel' e 'ilegivel'.")

    resultados = processar_pasta(pasta, usar_ia, threads)
    if not resultados:
        return

    grave = leve = acerto = revisar = 0
    for r in resultados:
        esperado = LEGIVEL if _normalizar(r["Colaborador"]) == "LEGIVEL" else ILEGIVEL
        r["Esperado"] = esperado
        if r["Status"] == REVISAR:
            r["Avaliação"] = "Foi para revisão"
            revisar += 1
        elif r["Status"] == esperado:
            r["Avaliação"] = "ACERTO"
            acerto += 1
        elif esperado == ILEGIVEL:
            r["Avaliação"] = "ERRO GRAVE: ilegível aprovado"
            grave += 1
        else:
            r["Avaliação"] = "Erro leve: legível reprovado"
            leve += 1

    total = len(resultados)
    decididos = total - revisar
    taxa = acerto / decididos if decididos else 0
    extras = {
        "Esperado x obtido — acertos": acerto,
        "ERROS GRAVES (ilegível aprovado)": grave,
        "Erros leves (legível reprovado)": leve,
        "Foram para revisão": revisar,
        "Taxa de acerto nas decisões automáticas": f"{taxa:.1%}",
        "Automação (% decidido sem humano)": f"{decididos / total:.1%}",
    }

    print("\n========== RESULTADO DA CALIBRAÇÃO ==========")
    for k, v in extras.items():
        print(f"{k:<45} {v}")

    print("\nNitidez (Laplaciano) por grupo — use para ajustar LIMITES['desfoque_alerta']:")
    for grupo in (LEGIVEL, ILEGIVEL):
        valores = [p["desfoque"] for r in resultados if r["Esperado"] == grupo
                   for p in r["_paginas"] if p["desfoque"] is not None]
        if valores:
            print(f"  {grupo:<9} mín {min(valores):8.1f} | mediana {float(np.median(valores)):8.1f}"
                  f" | máx {max(valores):8.1f}")

    destino = pasta / f"resultado_calibracao_{datetime.now():%Y-%m-%d_%H%M}.xlsx"
    for r in resultados:  # colunas extras na aba Documentos
        r.setdefault("Esperado", "")
    salvar_planilha(resultados, destino, extras)
    print(f"\nPlanilha salva em: {destino}")


# =============================================================================
# 9. PROGRAMA PRINCIPAL
# =============================================================================

def main() -> None:
    parser = argparse.ArgumentParser(description="Validador de legibilidade de documentos")
    parser.add_argument("pasta", nargs="?", help="Pasta com os documentos (subpasta por colaborador)")
    parser.add_argument("--calibrar", metavar="PASTA_TESTE",
                        help="Mede a taxa de acerto usando as subpastas 'legivel' e 'ilegivel'")
    parser.add_argument("--sem-ia", action="store_true", help="Roda só o filtro técnico (sem custo)")
    parser.add_argument("--colaboradores", metavar="PLANILHA",
                        help="Planilha com NOME e FUNÇÃO: confere o checklist de documentos e NRs/RACs")
    parser.add_argument("--threads", type=int, default=THREADS, help="Arquivos em paralelo")
    args = parser.parse_args()

    try:
        if args.calibrar:
            calibrar(Path(args.calibrar).resolve(), not args.sem_ia, args.threads)
            return
        if not args.pasta:
            parser.error("informe a pasta com os documentos")

        raiz = Path(args.pasta).resolve()
        if not raiz.is_dir():
            sys.exit(f"Pasta não encontrada: {raiz}")

        executar_analise(raiz, not args.sem_ia, args.threads,
                         planilha_colaboradores=Path(args.colaboradores) if args.colaboradores else None)
    except RuntimeError as erro:
        sys.exit(f"ERRO: {erro}")


if __name__ == "__main__":
    main()
